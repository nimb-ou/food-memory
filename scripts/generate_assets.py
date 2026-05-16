"""Generate lightweight portfolio assets from Food Memory result JSON.

The script deliberately keeps generated files outside source control by default:

    python scripts/generate_assets.py --results docs/sample_results.json --output docs/generated
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path


def load_results(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def pct(value: float) -> str:
    return f"{value * 100:.1f}%"


def generate_markdown(results: dict, out_dir: Path) -> Path:
    retrieval = results.get("retrieval", {})
    topk = retrieval.get("topk", {})
    latency = retrieval.get("latency_us", {})
    text = f"""# Food Memory Case Study

## Thesis

Food Memory applies retrieval-first machine learning to Food-101. It hashes exact images, embeds misses with CLIP, searches FAISS, and votes across nearest neighbors.

## Headline Metrics

| Metric | Value |
|---|---:|
| Retrieval top-1 | {pct(topk.get("top_1", 0.0))} |
| Retrieval top-3 | {pct(topk.get("top_3", 0.0))} |
| Retrieval top-5 | {pct(topk.get("top_5", 0.0))} |
| Macro F1 | {pct(retrieval.get("macro_f1", 0.0))} |
| p50 latency | {latency.get("p50", 0.0):.1f} us |
| p95 latency | {latency.get("p95", 0.0):.1f} us |

## Architecture

query image -> exact RGB hash -> CLIP embedding on miss -> FAISS search -> k-NN vote

## Baselines

The generated evaluation compares retrieval with CLIP zero-shot classification and logistic regression over frozen CLIP image embeddings.
"""
    path = out_dir / "food_memory_case_study.md"
    path.write_text(text, encoding="utf-8")
    return path


def generate_docx(results: dict, out_dir: Path) -> Path | None:
    try:
        from docx import Document
        from docx.shared import Inches, Pt
    except ImportError:
        return None

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)

    doc.add_heading("Food Memory Case Study", level=0)
    doc.add_paragraph("Retrieval-first image classification on Food-101 with exact memory and FAISS fallback.")
    doc.add_heading("Headline Metrics", level=1)
    retrieval = results.get("retrieval", {})
    topk = retrieval.get("topk", {})
    rows = [
        ("Retrieval top-1", pct(topk.get("top_1", 0.0))),
        ("Retrieval top-3", pct(topk.get("top_3", 0.0))),
        ("Retrieval top-5", pct(topk.get("top_5", 0.0))),
        ("Macro F1", pct(retrieval.get("macro_f1", 0.0))),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Metric"
    table.rows[0].cells[1].text = "Value"
    for name, value in rows:
        cells = table.add_row().cells
        cells[0].text = name
        cells[1].text = value
    doc.add_heading("Architecture", level=1)
    doc.add_paragraph("query image -> exact RGB hash -> CLIP embedding on miss -> FAISS search -> k-NN vote")
    doc.add_heading("Why It Matters", level=1)
    doc.add_paragraph(
        "The project establishes a measurable retrieval baseline before any fine-tuning. "
        "If memory is competitive, the system is fast, inspectable, and simple. "
        "If it is not, the nearest-neighbor failures reveal what a trained model must learn."
    )
    path = out_dir / "food_memory_case_study.docx"
    doc.save(path)
    return path


def generate_pptx(results: dict, out_dir: Path) -> Path | None:
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ImportError:
        return None

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    retrieval = results.get("retrieval", {})
    topk = retrieval.get("topk", {})

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.7), Inches(0.6), Inches(8.8), Inches(0.8))
    title.text_frame.text = "Food Memory"
    title.text_frame.paragraphs[0].font.size = Pt(44)
    subtitle = slide.shapes.add_textbox(Inches(0.75), Inches(1.45), Inches(9.5), Inches(0.55))
    subtitle.text_frame.text = "Retrieval-first image classification on Food-101"

    metric = slide.shapes.add_textbox(Inches(0.8), Inches(2.5), Inches(3.5), Inches(1.0))
    metric.text_frame.text = pct(topk.get("top_1", 0.0))
    metric.text_frame.paragraphs[0].font.size = Pt(48)
    label = slide.shapes.add_textbox(Inches(0.85), Inches(3.35), Inches(4.2), Inches(0.4))
    label.text_frame.text = "retrieval top-1 after evaluation"

    body = slide.shapes.add_textbox(Inches(5.2), Inches(2.4), Inches(6.8), Inches(2.0))
    body.text_frame.text = "Exact hash first. CLIP embedding on miss. FAISS search. k-NN vote. Baselines included."

    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    title2 = slide2.shapes.add_textbox(Inches(0.7), Inches(0.6), Inches(8.8), Inches(0.8))
    title2.text_frame.text = "Architecture"
    title2.text_frame.paragraphs[0].font.size = Pt(38)
    flow = slide2.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.5), Inches(1.2))
    flow.text_frame.text = "query image -> exact RGB hash -> CLIP embedding -> FAISS search -> k-NN vote"
    flow.text_frame.paragraphs[0].font.size = Pt(24)

    path = out_dir / "food_memory_deck.pptx"
    prs.save(path)
    return path


def generate_pdf(results: dict, out_dir: Path) -> Path | None:
    try:
        import matplotlib.pyplot as plt
        from matplotlib.backends.backend_pdf import PdfPages
    except ImportError:
        return None

    path = out_dir / "food_memory_poster.pdf"
    retrieval = results.get("retrieval", {})
    topk = retrieval.get("topk", {})
    with PdfPages(path) as pdf:
        fig = plt.figure(figsize=(11, 8.5))
        fig.patch.set_facecolor("#F2EFE6")
        ax = fig.add_axes([0, 0, 1, 1])
        ax.axis("off")
        ax.text(0.08, 0.82, "Food Memory", fontsize=42, weight="bold", color="#131312")
        ax.text(0.08, 0.73, "Retrieval-first image classification on Food-101", fontsize=17, color="#333333")
        ax.text(0.08, 0.52, pct(topk.get("top_1", 0.0)), fontsize=68, weight="bold", color="#B8331F")
        ax.text(0.09, 0.46, "retrieval top-1 after evaluation", fontsize=14, color="#333333")
        ax.text(0.08, 0.28, "hash -> CLIP -> FAISS -> k-NN vote", fontsize=24, color="#131312")
        ax.text(0.08, 0.17, f"Macro F1: {pct(retrieval.get('macro_f1', 0.0))}", fontsize=14, color="#333333")
        pdf.savefig(fig)
        plt.close(fig)
    return path


def parse_args() -> argparse.Namespace:
    ap = argparse.ArgumentParser(description="Generate Food Memory portfolio assets.")
    ap.add_argument("--results", type=Path, default=Path("docs/sample_results.json"))
    ap.add_argument("--output", type=Path, default=Path("docs/generated"))
    ap.add_argument("--formats", nargs="+", default=["md", "docx", "pptx", "pdf"], choices=["md", "docx", "pptx", "pdf"])
    return ap.parse_args()


def main() -> None:
    args = parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    results = load_results(args.results)
    created = []
    if "md" in args.formats:
        created.append(generate_markdown(results, args.output))
    if "docx" in args.formats:
        created.append(generate_docx(results, args.output))
    if "pptx" in args.formats:
        created.append(generate_pptx(results, args.output))
    if "pdf" in args.formats:
        created.append(generate_pdf(results, args.output))
    for path in created:
        if path is not None:
            print(path)


if __name__ == "__main__":
    main()

